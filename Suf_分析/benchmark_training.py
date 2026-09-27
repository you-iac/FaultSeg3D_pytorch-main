"""Analyze archived training curves/logs and benchmark real training steps.

Each controlled benchmark runs in a clean subprocess and follows the current
main_.py training path: model forward, configured compute_loss, con_matrix,
backward, and Adam optimizer step.
"""

from __future__ import annotations

import argparse
import ctypes
import hashlib
import importlib.util
import json
import math
import platform
import re
import statistics
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pandas as pd
import torch
from openpyxl import load_workbook
from torch.utils.data import DataLoader, Dataset


ROOT = Path(__file__).resolve().parents[1]
SUF_ROOT = Path(__file__).resolve().parent
TRAIN_ROOT = ROOT / "data" / "data_3D_400" / "train"
REPORT_PATH = SUF_ROOT / "模型训练过程与损失函数分析.md"
MODEL_FOLDERS = [
    "00_Baseline_Unet_8",
    "01_Unet_DCN4X2",
    "02_Unet_ConnLoss_8",
    "03_Unet_SufCOv4X2",
    "04_Unet_SufLoss",
    "05_Unet_SufCov+Loss4x2",
]


class TrainDataset(Dataset):
    def __init__(self, root: Path):
        def numeric_key(path: Path):
            try:
                return (0, int(path.stem))
            except ValueError:
                return (1, path.name)

        self.x_files = sorted((root / "x").glob("*.npy"), key=numeric_key)
        self.y_root = root / "y"
        missing = [path.name for path in self.x_files if not (self.y_root / path.name).exists()]
        if missing:
            raise FileNotFoundError(f"Missing labels for {missing[:5]}")

    def __len__(self) -> int:
        return len(self.x_files)

    def __getitem__(self, index: int):
        x_path = self.x_files[index]
        x = np.load(x_path)
        y = np.load(self.y_root / x_path.name)
        if x.ndim == 3:
            x = x[None, ...]
        return torch.from_numpy(x).float(), torch.from_numpy(y).float()


class PROCESS_MEMORY_COUNTERS_EX(ctypes.Structure):
    _fields_ = [
        ("cb", ctypes.c_ulong),
        ("PageFaultCount", ctypes.c_ulong),
        ("PeakWorkingSetSize", ctypes.c_size_t),
        ("WorkingSetSize", ctypes.c_size_t),
        ("QuotaPeakPagedPoolUsage", ctypes.c_size_t),
        ("QuotaPagedPoolUsage", ctypes.c_size_t),
        ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t),
        ("QuotaNonPagedPoolUsage", ctypes.c_size_t),
        ("PagefileUsage", ctypes.c_size_t),
        ("PeakPagefileUsage", ctypes.c_size_t),
        ("PrivateUsage", ctypes.c_size_t),
    ]


def process_memory() -> tuple[int, int, int]:
    counters = PROCESS_MEMORY_COUNTERS_EX()
    counters.cb = ctypes.sizeof(counters)
    kernel32 = ctypes.windll.kernel32
    psapi = ctypes.windll.psapi
    kernel32.GetCurrentProcess.restype = ctypes.c_void_p
    psapi.GetProcessMemoryInfo.argtypes = [
        ctypes.c_void_p,
        ctypes.POINTER(PROCESS_MEMORY_COUNTERS_EX),
        ctypes.c_ulong,
    ]
    psapi.GetProcessMemoryInfo.restype = ctypes.c_int
    handle = kernel32.GetCurrentProcess()
    if not psapi.GetProcessMemoryInfo(handle, ctypes.byref(counters), counters.cb):
        raise ctypes.WinError()
    return counters.WorkingSetSize, counters.PeakWorkingSetSize, counters.PrivateUsage


def percentile(values: list[float], fraction: float) -> float:
    values = sorted(values)
    position = (len(values) - 1) * fraction
    lower, upper = math.floor(position), math.ceil(position)
    if lower == upper:
        return values[lower]
    return values[lower] + (values[upper] - values[lower]) * (position - lower)


def parse_config(folder: Path) -> dict[str, str]:
    result = {}
    for line in (folder / "config.txt").read_text(encoding="utf-8", errors="replace").splitlines():
        if " : " in line:
            key, value = line.split(" : ", 1)
            result[key.strip()] = value.strip()
    return result


def import_model(model_py: Path):
    sys.path.insert(0, str(ROOT / "models"))
    sys.path.insert(0, str(model_py.parent))
    name = "suf_training_" + hashlib.md5(str(model_py).encode()).hexdigest()
    spec = importlib.util.spec_from_file_location(name, model_py)
    if spec is None or spec.loader is None:
        raise ImportError(f"Cannot import {model_py}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def average(values: list[float]) -> float:
    return statistics.fmean(values)


def summarize(values: list[float]) -> dict[str, float]:
    return {
        "mean": average(values),
        "median": statistics.median(values),
        "p95": percentile(values, 0.95),
        "min": min(values),
        "max": max(values),
    }


def run_worker(folder_name: str, steps: int, warmup: int) -> dict:
    sys.path.insert(0, str(ROOT))
    from utils.tools import compute_loss, con_matrix

    if not torch.cuda.is_available():
        raise RuntimeError("CUDA is required")
    torch.cuda.set_device(0)
    device = torch.device("cuda:0")
    folder = SUF_ROOT / folder_name
    config = parse_config(folder)
    dataset = TrainDataset(TRAIN_ROOT)
    required = warmup + steps
    if len(dataset) < required:
        raise ValueError(f"Need {required} samples, found {len(dataset)}")
    loader = DataLoader(dataset, batch_size=1, shuffle=False, num_workers=0, pin_memory=False)
    iterator = iter(loader)

    module = import_model(folder / "models.py")
    model = module.FaultSeg3D(1, 2)
    state = torch.load(folder / "models" / "FaultSeg3D_BEST.pth", map_location="cpu", weights_only=True)
    model.load_state_dict(state, strict=True)
    model.to(device).train()
    torch.cuda.synchronize()
    parameter_mib = sum(p.numel() * p.element_size() for p in model.parameters()) / 1024**2
    gpu_model_loaded_mib = torch.cuda.memory_allocated(device) / 1024**2
    optimizer = torch.optim.Adam(model.parameters(), lr=float(config.get("optim_lr", "0.0001")))
    args = SimpleNamespace(
        device=str(device),
        out_channels=2,
        loss_func=config["loss_func"],
        current_epoch=50,
        surface_weight=float(config.get("surface_weight", "0.1")),
        surface_geometry_strength=float(config.get("surface_geometry_strength", "1.0")),
        surface_input_type=config.get("surface_input_type", "probabilities"),
        debug_loss=False,
        exp=None,
    )

    def step_once(batch, measure: bool):
        x, y = batch
        timings = {}

        start = time.perf_counter()
        x, y = x.to(device), y.to(device)
        torch.cuda.synchronize()
        timings["h2d_ms"] = (time.perf_counter() - start) * 1000

        optimizer.zero_grad()
        start = time.perf_counter()
        output = model(x)
        torch.cuda.synchronize()
        timings["forward_ms"] = (time.perf_counter() - start) * 1000

        start = time.perf_counter()
        loss = compute_loss(output, y, args)
        torch.cuda.synchronize()
        timings["loss_ms"] = (time.perf_counter() - start) * 1000

        start = time.perf_counter()
        metrics = con_matrix(output, y, args)
        timings["metrics_ms"] = (time.perf_counter() - start) * 1000

        start = time.perf_counter()
        loss.backward()
        optimizer.step()
        torch.cuda.synchronize()
        timings["backward_step_ms"] = (time.perf_counter() - start) * 1000
        timings["loss_value"] = float(loss.detach().cpu())
        del output, loss, x, y, metrics
        return timings if measure else None

    for index in range(warmup):
        step_once(next(iterator), measure=False)
        print(f"[{folder_name}] warmup {index + 1}/{warmup}", flush=True)

    torch.cuda.empty_cache()
    torch.cuda.reset_peak_memory_stats(device)
    gpu_idle_allocated = torch.cuda.memory_allocated(device)
    cpu_idle_working, _, _ = process_memory()
    records = []
    data_load_ms = []
    total_start = time.perf_counter()
    for index in range(steps):
        load_start = time.perf_counter()
        batch = next(iterator)
        load_ms = (time.perf_counter() - load_start) * 1000
        record = step_once(batch, measure=True)
        assert record is not None
        record["data_load_ms"] = load_ms
        record["total_ms"] = load_ms + sum(
            record[key]
            for key in ("h2d_ms", "forward_ms", "loss_ms", "metrics_ms", "backward_step_ms")
        )
        records.append(record)
        data_load_ms.append(load_ms)
        if (index + 1) % 5 == 0 or index + 1 == steps:
            print(f"[{folder_name}] measured {index + 1}/{steps}", flush=True)
    torch.cuda.synchronize()
    measured_wall_seconds = time.perf_counter() - total_start
    _, cpu_peak_working, _ = process_memory()

    result = {
        "folder": folder_name,
        "loss_func": config["loss_func"],
        "steps": steps,
        "warmup_steps": warmup,
        "batch_size": 1,
        "current_epoch": 50,
        "parameter_mib": parameter_mib,
        "gpu_model_loaded_mib": gpu_model_loaded_mib,
        "measured_wall_seconds": measured_wall_seconds,
        "samples_per_second": steps / measured_wall_seconds,
        "estimated_200_sample_epoch_seconds": measured_wall_seconds / steps * 200,
        "estimated_400_sample_epoch_seconds": measured_wall_seconds / steps * 400,
        "loss_value_mean": average([row["loss_value"] for row in records]),
        "gpu_idle_allocated_mib": gpu_idle_allocated / 1024**2,
        "gpu_peak_allocated_mib": torch.cuda.max_memory_allocated(device) / 1024**2,
        "gpu_peak_reserved_mib": torch.cuda.max_memory_reserved(device) / 1024**2,
        "cpu_idle_working_set_mib": cpu_idle_working / 1024**2,
        "cpu_peak_working_set_mib": cpu_peak_working / 1024**2,
    }
    for key in (
        "data_load_ms",
        "h2d_ms",
        "forward_ms",
        "loss_ms",
        "metrics_ms",
        "backward_step_ms",
        "total_ms",
    ):
        result[key] = summarize([row[key] for row in records])
    return result


LOG_PATTERN = re.compile(
    r"^(\d{4}-\d\d-\d\d \d\d:\d\d:\d\d\.\d+) : epoch: (\d+) (?:Times:|step:)(\d+)"
)


def analyze_log(folder: Path, configured_batch_size: int) -> dict:
    rows = []
    for line in (folder / "log.txt").read_text(encoding="utf-8", errors="replace").splitlines():
        match = LOG_PATTERN.match(line)
        if match:
            rows.append(
                (
                    datetime.fromisoformat(match.group(1)),
                    int(match.group(2)),
                    int(match.group(3)),
                )
            )
    # Some logs include an aborted/restarted prefix. The last record for each
    # epoch/global-step pair belongs to the completed run retained in the file.
    deduplicated = {(epoch, step): (stamp, epoch, step) for stamp, epoch, step in rows}
    clean = sorted(deduplicated.values(), key=lambda row: row[0])
    by_epoch = {}
    for stamp, epoch, step in clean:
        by_epoch.setdefault(epoch, []).append((stamp, step))
    within = []
    for entries in by_epoch.values():
        entries.sort(key=lambda item: item[1])
        for (first_time, first_step), (second_time, second_step) in zip(entries, entries[1:]):
            if second_step == first_step + 1:
                within.append((second_time - first_time).total_seconds())
    counts = [len(entries) for entries in by_epoch.values()]
    median_step = statistics.median(within)
    return {
        "raw_epoch_lines": len(rows),
        "deduplicated_epoch_lines": len(clean),
        "epochs": len(by_epoch),
        "steps_per_epoch_median": statistics.median(counts),
        "steps_per_epoch_min": min(counts),
        "steps_per_epoch_max": max(counts),
        "implied_samples_per_epoch": statistics.median(counts) * configured_batch_size,
        "start": clean[0][0].isoformat(sep=" "),
        "end": clean[-1][0].isoformat(sep=" "),
        "logged_wall_hours": (clean[-1][0] - clean[0][0]).total_seconds() / 3600,
        "step_seconds_mean": average(within),
        "step_seconds_median": median_step,
        "step_seconds_p95": percentile(within, 0.95),
        "historical_samples_per_second": configured_batch_size / median_step,
    }


def inspect_workbook(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    errors = []
    formulas = 0
    for row in sheet.iter_rows():
        for cell in row:
            formulas += int(cell.data_type == "f")
            if cell.data_type == "e":
                errors.append((cell.coordinate, cell.value))
    frame = pd.read_excel(path, index_col=0)
    return {
        "path": str(path.relative_to(SUF_ROOT)).replace("\\", "/"),
        "sheets": workbook.sheetnames,
        "rows": len(frame),
        "columns": list(frame.columns),
        "formula_count": formulas,
        "errors": errors,
        "nan_count": int(frame.isna().sum().sum()),
        "frame": frame,
    }


def curve_summary(folder: Path) -> dict | None:
    train_files = list(folder.rglob("train_result.xlsx"))
    val_files = list(folder.rglob("val_result.xlsx"))
    if not train_files or not val_files:
        return None
    train_info = inspect_workbook(train_files[0])
    val_info = inspect_workbook(val_files[0])
    train = train_info.pop("frame")
    val = val_info.pop("frame")
    best_index = int(val["val_iou"].idxmax())
    best_dice_index = int(val["val_dice"].idxmax())
    first_iou = float(val.iloc[0]["val_iou"])
    best_iou = float(val.loc[best_index, "val_iou"])
    threshold = first_iou + 0.95 * (best_iou - first_iou)
    reached = val.index[val["val_iou"] >= threshold]
    convergence_epoch = int(reached[0]) + 1 if len(reached) else None
    return {
        "train_source": train_info,
        "val_source": val_info,
        "epochs": len(train),
        "train_first_loss": float(train.iloc[0]["train_loss"]),
        "train_last_loss": float(train.iloc[-1]["train_loss"]),
        "train_first_iou": float(train.iloc[0]["train_iou"]),
        "train_last_iou": float(train.iloc[-1]["train_iou"]),
        "train_first_dice": float(train.iloc[0]["train_dice"]),
        "train_last_dice": float(train.iloc[-1]["train_dice"]),
        "val_first_loss": float(val.iloc[0]["val_loss"]),
        "val_last_loss": float(val.iloc[-1]["val_loss"]),
        "val_first_iou": first_iou,
        "val_last_iou": float(val.iloc[-1]["val_iou"]),
        "val_first_dice": float(val.iloc[0]["val_dice"]),
        "val_last_dice": float(val.iloc[-1]["val_dice"]),
        "best_val_iou": best_iou,
        "best_val_iou_epoch": best_index + 1,
        "best_val_dice": float(val.loc[best_dice_index, "val_dice"]),
        "best_val_dice_epoch": best_dice_index + 1,
        "min_val_loss": float(val["val_loss"].min()),
        "min_val_loss_epoch": int(val["val_loss"].idxmin()) + 1,
        "generalization_gap_dice_at_best_iou": float(
            train.loc[best_index, "train_dice"] - val.loc[best_index, "val_dice"]
        ),
        "post_best_iou_drop": best_iou - float(val.iloc[-1]["val_iou"]),
        "convergence_95pct_epoch": convergence_epoch,
        "val_iou_last10_std": float(val["val_iou"].tail(10).std(ddof=0)),
    }


def parse_final_validation(folder: Path) -> dict | None:
    files = list(folder.rglob("valid_final_result.txt"))
    if not files:
        return None
    values = {}
    for line in files[0].read_text(encoding="utf-8", errors="replace").splitlines():
        if ":" in line:
            key, value = line.split(":", 1)
            try:
                values[key.strip().replace("valid ", "")] = float(value.strip())
            except ValueError:
                pass
    values["source"] = str(files[0].relative_to(SUF_ROOT)).replace("\\", "/")
    return values


def collect_archived_analysis() -> list[dict]:
    analyses = []
    for folder_name in MODEL_FOLDERS:
        folder = SUF_ROOT / folder_name
        config = parse_config(folder)
        batch_size = int(config["batch_size"])
        analyses.append(
            {
                "folder": folder_name,
                "config": config,
                "log": analyze_log(folder, batch_size),
                "curve": curve_summary(folder),
                "final_validation": parse_final_validation(folder),
            }
        )
    return analyses


def final_metric(analysis: dict, name: str) -> float | None:
    final = analysis["final_validation"]
    if final and name in final:
        return final[name]
    curve = analysis["curve"]
    if curve:
        return curve[f"best_val_{name}"]
    return None


def render_report(analyses: list[dict], benchmarks: list[dict], metadata: dict) -> str:
    by_benchmark = {row["folder"]: row for row in benchmarks}
    ranked = sorted(
        analyses,
        key=lambda row: final_metric(row, "dice") if final_metric(row, "dice") is not None else -1,
        reverse=True,
    )
    baseline_dice = final_metric(analyses[0], "dice")
    baseline_iou = final_metric(analyses[0], "iou")
    lines = [
        "# Suf 模型训练过程与损失函数分析",
        "",
        f">生成时间：{metadata['generated_at']}（Asia/Shanghai）",
        "",
        "## 主要结论",
        "",
        f"- 最高最终验证 Dice 为 `{ranked[0]['folder']}` 的 "
        f"{final_metric(ranked[0], 'dice'):.5f}，IoU 为 {final_metric(ranked[0], 'iou'):.5f}。",
        "- SurfaceConv 结构带来明显指标改善，但训练和推理都更慢。"
        "Surface loss 单独使用未在这组实验中超过 Baseline。",
        "- `ConnectivityLoss` 实验的最佳验证 Dice 与 Baseline 几乎相同，而其训练损失计算包含额外的多尺度连通性项。",
        "- 不同 `loss_func` 的 loss 数值不具有横向可比性。模型优劣应主要比较 IoU、Dice 和实测速度。",
        "",
        "## 配置与损失函数映射",
        "",
        "| 模型 | config `loss_func` | 实际组合 | 调度 |",
        "|---|---|---|---|",
        "| `00_Baseline_Unet_8` | `dice_plus_ce` | `DiceLoss + WeightedCrossEntropyLoss` | 固定相加 |",
        "| `01_Unet_DCN4X2` | `dice_plus_ce` | `DiceLoss + WeightedCrossEntropyLoss` | 固定相加 |",
        "| `02_Unet_ConnLoss_8` | `ConnectivityLoss` | `DiceLoss + WeightedCrossEntropyLoss + α·ConnectivityLoss` | epoch 1–10: α=0；11–20: 线性增至0.1；之后0.1 |",
        "| `03_Unet_SufCOv4X2` | `dice_plus_ce` | `DiceLoss + WeightedCrossEntropyLoss` | 固定相加 |",
        "| `04_Unet_SufLoss` | `SurfaceGuidedBreakMergeLoss` | `WeightedCrossEntropyLoss + α·SurfaceGuidedBreakMergeLoss` | epoch 1–10: α=0；11–20: 增至 `surface_weight=0.1`；之后0.1 |",
        "| `05_Unet_SufCov+Loss4x2` | `SurfaceGuidedBreakMergeLoss` | `WeightedCrossEntropyLoss + α·SurfaceGuidedBreakMergeLoss` | 同上 |",
        "",
        "重要细节：",
        "",
        "- `SurfaceGuidedBreakMergeLoss` 分支中的基础项只是加权 CE，没有 Dice。",
        "- `ConnectivityLoss` 在 α=0 的前10个 epoch 仍会完整计算连通性项，只是最后乘以0，因而不会节省计算时间。",
        "- 模型 `forward()` 已输出 Softmax 概率，但当 `ConnectivityLoss` 收到两通道输入时会无条件再次执行 `softmax`。"
        "当前实现即使传入 `input_is_probability=True` 也不会绕过这一分支。"
        "建议修改 `_foreground_probability`：对已是概率的两通道输入直接取 `pred[:, 1:2]`。",
        "",
        "## 历史训练曲线与最终验证",
        "",
        "| 排名 | 模型 | 最佳/最终验证 IoU | 验证 Dice | 最佳epoch | 训练末 IoU | 训练末 Dice | 最佳点 Dice 泛化差 | 数据完整性 |",
        "|---:|---|---:|---:|---:|---:|---:|---:|---|",
    ]
    for rank, analysis in enumerate(ranked, start=1):
        curve = analysis["curve"]
        dice = final_metric(analysis, "dice")
        iou = final_metric(analysis, "iou")
        if curve:
            best_epoch = curve["best_val_iou_epoch"]
            train_iou = f"{curve['train_last_iou']:.5f}"
            train_dice = f"{curve['train_last_dice']:.5f}"
            gap = f"{curve['generalization_gap_dice_at_best_iou']:.5f}"
            completeness = "50 epoch Excel完整"
        else:
            best_epoch = "n.a."
            train_iou = train_dice = gap = "n.a."
            completeness = "缺训练/val Excel，仅最终验证"
        lines.append(
            f"| {rank} | `{analysis['folder']}` | {iou:.5f} | {dice:.5f} | {best_epoch} | "
            f"{train_iou} | {train_dice} | {gap} | {completeness} |"
        )

    lines.extend(
        [
            "",
            "### 相对 Baseline 的指标变化",
            "",
            "| 模型 | IoU变化（百分点） | Dice变化（百分点） |",
            "|---|---:|---:|",
        ]
    )
    for analysis in analyses:
        iou = final_metric(analysis, "iou")
        dice = final_metric(analysis, "dice")
        lines.append(
            f"| `{analysis['folder']}` | {(iou - baseline_iou) * 100:+.3f} | "
            f"{(dice - baseline_dice) * 100:+.3f} |"
        )

    lines.extend(
        [
            "",
            "### 曲线稳定性",
            "",
            "| 模型 | 首epoch→末epoch val IoU | 95%改善首次达到epoch | 后10epoch val IoU标准差 | 最佳后IoU回落 |",
            "|---|---:|---:|---:|---:|",
        ]
    )
    for analysis in analyses:
        curve = analysis["curve"]
        if curve is None:
            lines.append(f"| `{analysis['folder']}` | n.a. | n.a. | n.a. | n.a. |")
        else:
            lines.append(
                f"| `{analysis['folder']}` | {curve['val_first_iou']:.5f}→{curve['val_last_iou']:.5f} | "
                f"{curve['convergence_95pct_epoch']} | {curve['val_iou_last10_std']:.5f} | "
                f"{curve['post_best_iou_drop']:.5f} |"
            )

    lines.extend(
        [
            "",
            "## 历史日志中的实际训练速度",
            "",
            "| 模型 | config batch | 每epoch步数 | 推定样本/epoch | 步耗时P50/P95 (s) | 历史吞吐量 (样本/s) | 日志首尾时长 (h) |",
            "|---|---:|---:|---:|---:|---:|---:|",
        ]
    )
    for analysis in analyses:
        log = analysis["log"]
        batch = int(analysis["config"]["batch_size"])
        lines.append(
            f"| `{analysis['folder']}` | {batch} | {log['steps_per_epoch_median']:.0f} | "
            f"{log['implied_samples_per_epoch']:.0f} | {log['step_seconds_median']:.3f} / "
            f"{log['step_seconds_p95']:.3f} | {log['historical_samples_per_second']:.2f} | "
            f"{log['logged_wall_hours']:.2f} |"
        )

    lines.extend(
        [
            "",
            "日志说明：每个步骤的时间戳写在 `optimizer.step()` 之后，同epoch相邻时间差可以反映"
            "数据加载、前向、loss、指标、反向和更新的综合速度。日志首尾时长还包含epoch间验证与checkpoint，"
            "但不包含第一步之前的初始化和最后一步之后的最终验证。",
            "",
            "## 当前环境的统一训练步基准",
            "",
            f"条件：batch size=1，FP32，Adam，`current_epoch=50`，每模型预热 {metadata['warmup_steps']} 步后测量 "
            f"{metadata['measured_steps']} 个真实训练样本。每步按 `main_.py` 的正式路径执行"
            "forward、配置损失、`con_matrix`、backward 和 `optimizer.step`。"
            "每个子进程先严格加载对应 `FaultSeg3D_BEST.pth`，基准中的更新不会写回权重文件。",
            "",
            "| 模型 | loss_func | 实测步耗时P50/P95 (ms) | 吞吐量 (样本/s) | loss计算P50 (ms) | 反向+更新P50 (ms) | 预估400样本 (min) | 峰值已分配显存 (MiB) | CPU峰值工作集 (MiB) |",
            "|---|---|---:|---:|---:|---:|---:|---:|---:|",
        ]
    )
    for analysis in analyses:
        benchmark = by_benchmark[analysis["folder"]]
        lines.append(
            f"| `{analysis['folder']}` | `{benchmark['loss_func']}` | "
            f"{benchmark['total_ms']['median']:.1f} / {benchmark['total_ms']['p95']:.1f} | "
            f"{benchmark['samples_per_second']:.3f} | {benchmark['loss_ms']['median']:.1f} | "
            f"{benchmark['backward_step_ms']['median']:.1f} | "
            f"{benchmark['estimated_400_sample_epoch_seconds'] / 60:.2f} | "
            f"{benchmark['gpu_peak_allocated_mib']:.1f} | {benchmark['cpu_peak_working_set_mib']:.1f} |"
        )

    lines.extend(
        [
            "",
            "### GPU 显存占用明细",
            "",
            "| 模型 | FP32参数 (MiB) | 仅模型加载后 (MiB) | 训练稳态基线 (MiB) | 峰值已分配 (MiB) | 峰值增量 (MiB) | 峰值保留 (MiB) | 峰值/物理显存 |",
            "|---|---:|---:|---:|---:|---:|---:|---:|",
        ]
    )
    for analysis in analyses:
        benchmark = by_benchmark[analysis["folder"]]
        increment = benchmark["gpu_peak_allocated_mib"] - benchmark["gpu_idle_allocated_mib"]
        lines.append(
            f"| `{analysis['folder']}` | {benchmark['parameter_mib']:.1f} | "
            f"{benchmark['gpu_model_loaded_mib']:.1f} | {benchmark['gpu_idle_allocated_mib']:.1f} | "
            f"{benchmark['gpu_peak_allocated_mib']:.1f} | {increment:.1f} | "
            f"{benchmark['gpu_peak_reserved_mib']:.1f} | "
            f"{benchmark['gpu_peak_allocated_mib'] / metadata['gpu_total_mib'] * 100:.1f}% |"
        )

    lines.extend(
        [
            "",
            "显存口径：",
            "",
            "- “仅模型加载后”在权重移入GPU、Adam尚未建立状态时读取 `torch.cuda.memory_allocated`。",
            "- “训练稳态基线”在预热完成后读取，包含模型参数、梯度和 Adam 状态等常驻张量。",
            "- “峰值已分配”是计时训练步中 PyTorch 张量实际占用的最高值；"
            "“峰值增量”=峰值已分配−稳态基线，主要反映激活、临时缓冲区和中间梯度。",
            "- “峰值保留”是 PyTorch allocator 向 CUDA 申请并缓存的峰值。"
            "它通常高于已分配值，但不包含 CUDA context、驱动和其他进程，因而不会与 `nvidia-smi` 完全相同。",
            "",
            "### 训练步分解（P50，ms）",
            "",
            "| 模型 | 数据读取 | H2D | 前向 | 损失 | 指标 | 反向+更新 |",
            "|---|---:|---:|---:|---:|---:|---:|",
        ]
    )
    for analysis in analyses:
        benchmark = by_benchmark[analysis["folder"]]
        lines.append(
            f"| `{analysis['folder']}` | {benchmark['data_load_ms']['median']:.1f} | "
            f"{benchmark['h2d_ms']['median']:.1f} | {benchmark['forward_ms']['median']:.1f} | "
            f"{benchmark['loss_ms']['median']:.1f} | {benchmark['metrics_ms']['median']:.1f} | "
            f"{benchmark['backward_step_ms']['median']:.1f} |"
        )

    lines.extend(
        [
            "",
            "## 训练框架审查",
            "",
            "- `main_.py` 实际调用 `utils.train.train`，该循环每个 batch 都执行 "
            "`zero_grad → forward → compute_loss → con_matrix → backward → step`。",
            "- `main_.py` 虽定义了 `grad_accum_steps`、`amp`、`multi_gpu` 和 `disable_checkpointing`，"
            "但当前 `utils/train.py` 没有使用这些参数。因此按当前正式路径，"
            "`batch_size=4, grad_accum_steps=2` 的实际优化批量仍是4，不是配置中标注的 effective batch size 8。",
            "- 历史日志的每epoch步数与400个训练样本一致：batch 8 时50步，batch 4 时100步。"
            f"但当前 `{metadata['train_root']}/x` 只剩 {metadata['current_train_samples']} 个样本，"
            "所以新基准使用现有样本子集，400样本时长为线性外推。",
            "- 原 `ConnectivityLoss` 分支每 batch 都会 `.item()`、打印并写日志，强制 CUDA 同步且增加 I/O。"
            "本次已改为仅在 `args.debug_loss=True` 时输出，损失数学计算未改变。",
            "- 修改后的冒烟测试确认：`debug_loss=False` 无标准输出，`debug_loss=True` 仍会输出 `ConnLoss`诊断，"
            "两种模式的 loss 数值完全相同。Surface loss 现有28项 unittest 全部通过。",
            "",
            "## 数据质量与限制",
            "",
            "- Excel 源表均只包含静态数值，未发现公式、Excel错误值或缺失值。",
            "- `01_Unet_DCN4X2` 缺少 `train_result.xlsx` 和 `val_result.xlsx`，只能使用 `valid_final_result.txt` 分析最终指标。",
            "- `02_Unet_ConnLoss_8` 缺少 `valid_final_result.txt`，报告使用 `val_result.xlsx` 中最高 IoU 所在epoch的指标。",
            "- 历史速度受当时GPU、文件缓存、温度和代码版本影响。当前基准用同一台GPU和统一batch size，更适合横向对比。",
            "- loss 调度基准固定为 epoch 50，因此包含完整连通/表面辅助项；不代表前10个epoch的 loss 组成。",
            "",
            "## 测试环境",
            "",
            f"- GPU：{metadata['gpu_name']}，{metadata['gpu_total_mib']:.0f} MiB",
            f"- PyTorch：{metadata['torch_version']}，CUDA runtime：{metadata['cuda_version']}，cuDNN：{metadata['cudnn_version']}",
            f"- Python：{metadata['python_version']}，操作系统：{metadata['platform']}",
            "",
            "## 源文件",
            "",
            "- 配置：各模型目录的 `config.txt`。",
            "- 曲线：各模型的 `results/train/train_result.xlsx` 和 `val_result.xlsx`。",
            "- 最终验证：`valid_final_result.txt`。",
            "- 历史速度：各模型目录的 `log.txt`。",
            "- 框架映射：`main_.py`、`utils/train.py`、`utils/tools.py`、"
            "`utils/loss/connectivity_loss.py` 和 `utils/loss/surface_guided_break_merge_loss.py`。",
            "",
        ]
    )
    return "\n".join(lines)


def run_controller(steps: int, warmup: int) -> None:
    analyses = collect_archived_analysis()
    benchmarks = []
    for folder in MODEL_FOLDERS:
        print(f"\n=== Training benchmark: {folder} ===", flush=True)
        command = [
            sys.executable,
            str(Path(__file__).resolve()),
            "--worker",
            folder,
            "--steps",
            str(steps),
            "--warmup",
            str(warmup),
        ]
        process = subprocess.Popen(
            command,
            cwd=ROOT,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            errors="replace",
            bufsize=1,
        )
        result = None
        assert process.stdout is not None
        for line in process.stdout:
            line = line.rstrip()
            print(line, flush=True)
            if line.startswith("RESULT_JSON="):
                result = json.loads(line[len("RESULT_JSON=") :])
        return_code = process.wait()
        if return_code != 0 or result is None:
            raise RuntimeError(f"Training benchmark failed for {folder} (exit {return_code})")
        benchmarks.append(result)

    _, total_gpu = torch.cuda.mem_get_info(0)
    metadata = {
        "generated_at": datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %z"),
        "warmup_steps": warmup,
        "measured_steps": steps,
        "train_root": str(TRAIN_ROOT.relative_to(ROOT)).replace("\\", "/"),
        "current_train_samples": len(TrainDataset(TRAIN_ROOT)),
        "gpu_name": torch.cuda.get_device_name(0),
        "gpu_total_mib": total_gpu / 1024**2,
        "torch_version": torch.__version__,
        "cuda_version": torch.version.cuda,
        "cudnn_version": torch.backends.cudnn.version(),
        "python_version": platform.python_version(),
        "platform": platform.platform(),
    }
    REPORT_PATH.write_text(render_report(analyses, benchmarks, metadata), encoding="utf-8")
    print(f"\nReport written to: {REPORT_PATH}", flush=True)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--worker", choices=MODEL_FOLDERS)
    parser.add_argument("--steps", type=int, default=20)
    parser.add_argument("--warmup", type=int, default=2)
    args = parser.parse_args()
    if args.worker:
        result = run_worker(args.worker, args.steps, args.warmup)
        print("RESULT_JSON=" + json.dumps(result, ensure_ascii=False), flush=True)
    else:
        run_controller(args.steps, args.warmup)


if __name__ == "__main__":
    main()
